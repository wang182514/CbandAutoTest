"""
Insert bookmark names into CbandTemplate.docx at current hardcoded cell positions.
This creates a marked-up template showing where each bookmark should go.
Save as CbandTemplate_bookmarked.docx — use it as reference when manually adding bookmarks.
"""
import os
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

base_dir = os.path.dirname(os.path.abspath(__file__))
template_path = os.path.join(base_dir, "..", "CbandTemplate.docx")
if not os.path.exists(template_path):
    print(f"模板未找到: {template_path}")
    exit(1)

doc = Document(template_path)

# ── Current cell mapping (from utils/report.py) ──
# (table_index, row, col, current_value_source)
# We write the BOOKMARK NAME into each cell so you know what to insert.
cells = [
    # 产品信息 (table 1)
    (1, 1, 3, "product_name"),
    (1, 1, 5, "product_model"),
    (1, 2, 3, "test_date"),
    (1, 2, 5, "test_env"),
    (1, 3, 3, "serial_number"),
    (1, 3, 5, "operator"),
    # 接收 (table 2)
    (2, 1,  6, "rx_current"),
    (2, 4,  3, "gain_mean_db"),
    (2, 5,  3, "nf_mean_db"),
    (2, 6,  3, "gain_flatness_db"),
    (2, 7,  3, "rx_nf_flatness_limit"),
    (2, 8,  3, "rx_pn_spots"),
    # 发射 (table 2)
    (2, 9,  6, "tx_peak_current"),
    (2, 11, 5, "tx_gain_1050"),
    (2, 12, 5, "tx_pout_1050"),
    (2, 13, 5, "tx_gain_1200"),
    (2, 14, 5, "tx_pout_1200"),
    (2, 15, 5, "tx_gain_1550"),
    (2, 16, 5, "tx_pout_1550"),
    (2, 17, 3, "tx_gain_limit_note"),
    (2, 18, 3, "tx_flatness_db"),
    (2, 19, 3, "tx_flatness_limit"),
    (2, 20, 3, "tx_pn_spots"),
]

printed = set()
for tab_idx, row, col, bookmark_name in cells:
    try:
        tbl = doc.tables[tab_idx - 1]
        cell = tbl.cell(row - 1, col - 1)

        # Write the bookmark name as visible text
        cell.text = f"«{bookmark_name}»"

        if bookmark_name not in printed:
            print(f"  [{bookmark_name}] → 表格{tab_idx}, 行{row}, 列{col}")
            printed.add(bookmark_name)
    except Exception as e:
        print(f"  ✗ [{bookmark_name}] 无法写入表格{tab_idx}({row},{col}): {e}")

out_path = os.path.join(os.path.dirname(template_path), "CbandTemplate_bookmarked.docx")
doc.save(out_path)
print(f"\n已保存: {out_path}")
print("打开此文件，每个单元格内有 «书签名»，对照在 Word 里插入书签即可。")
